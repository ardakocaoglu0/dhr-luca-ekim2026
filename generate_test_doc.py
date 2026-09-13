from docx import Document
from docx.oxml import parse_xml
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

_CB_ID = 1000


def next_cb_id():
    global _CB_ID
    _CB_ID += 1
    return _CB_ID


def set_narrow_margins(section):
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)


def set_keep_with_next(paragraph, keep=True):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ('keepNext', 'keepLines'):
        existing = pPr.find(qn(f'w:{tag}'))
        if existing is not None:
            pPr.remove(existing)
        if keep:
            pPr.append(OxmlElement(f'w:{tag}'))


def set_page_break_before(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn('w:pageBreakBefore'))
    if existing is not None:
        pPr.remove(existing)
    pPr.append(OxmlElement('w:pageBreakBefore'))


def add_checkbox_sdt(paragraph):
    """Interactive Word checkbox (clickable in Word / Word Online)."""
    cb_id = next_cb_id()
    xml = (
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w:sdtPr>'
        f'<w:id w:val="{cb_id}"/>'
        '<w14:checkbox>'
        '<w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="Segoe UI Symbol"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="Segoe UI Symbol"/>'
        '</w14:checkbox>'
        '</w:sdtPr>'
        '<w:sdtContent>'
        '<w:r>'
        '<w:rPr>'
        '<w:rFonts w:ascii="Segoe UI Symbol" w:hAnsi="Segoe UI Symbol"/>'
        '<w:sz w:val="22"/>'
        '</w:rPr>'
        '<w:t>&#9744;</w:t>'
        '</w:r>'
        '</w:sdtContent>'
        '</w:sdt>'
    )
    paragraph._p.append(parse_xml(xml))


def add_check_item(doc, text, keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.05
    pf.left_indent = Cm(0.25)
    add_checkbox_sdt(p)
    run = p.add_run('  ' + text)
    run.font.name = 'Calibri'
    run.font.size = Pt(9.5)
    if keep_with_next:
        set_keep_with_next(p, True)
    return p


def add_section_heading(doc, text, page_break=False):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    set_keep_with_next(h, True)
    if page_break:
        set_page_break_before(h)
    for run in h.runs:
        run.font.color.rgb = RGBColor(20, 60, 110)
        run.font.size = Pt(13)
    return h


def add_subheading(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(7)
    h.paragraph_format.space_after = Pt(2)
    set_keep_with_next(h, True)
    for run in h.runs:
        run.font.color.rgb = RGBColor(40, 90, 140)
        run.font.size = Pt(11)
    return h


def add_h3(doc, text):
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(5)
    h.paragraph_format.space_after = Pt(2)
    set_keep_with_next(h, True)
    for run in h.runs:
        run.font.color.rgb = RGBColor(60, 100, 140)
        run.font.size = Pt(10.5)
    return h


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    set_keep_with_next(p, True)
    run = p.add_run(text)
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(90, 90, 90)
    return p


def add_bullets(doc, items):
    for i, item in enumerate(items):
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(item)
        run.font.size = Pt(9.5)
        if i < len(items) - 1:
            set_keep_with_next(p, True)


def add_check_group(doc, items):
    for i, text in enumerate(items):
        keep = i < min(2, len(items) - 1)
        add_check_item(doc, text, keep_with_next=keep)


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(2)

for section in doc.sections:
    set_narrow_margins(section)

title = doc.add_heading('DHR Release Notes & Test Planı', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
for run in title.runs:
    run.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run = p.add_run(
    f'Tarih: 01–04 Eylül 2026  |  Hazırlanma: {datetime.date.today().strftime("%d.%m.%Y")}'
)
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(100, 100, 100)

note = doc.add_paragraph()
note.paragraph_format.space_after = Pt(6)
run = note.add_run(
    'Kutulara tıklayarak işaretleyin (Word / Word Online). Ana bölümler yeni sayfada başlar (01–04 Eylül 2026, 31 bölüm).'
)
run.font.size = Pt(8)
run.font.italic = True
run.font.color.rgb = RGBColor(120, 120, 120)

# 1
add_section_heading(doc, '1. İzin ve Ek Mesai — Kurucu Yönetici Muafiyeti Kaldırıldı')
add_meta(doc, 'Ekran: İzin Yönetimi → Talepler · Ek Mesai Yönetimi → Talepler · Çalışan Kartı İzinler / Ek Mesailer')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Kurucu yönetici (admin@d1-tech.com) dahil hiçbir hesabın pozisyon muafiyeti kalmadı.',
    'Onay adımını yalnızca o adımın pozisyonundaki kullanıcı onaylayabilir / reddedebilir.',
    'Listeleme muafiyeti devam ediyor — kurucu yönetici tüm talepleri görebiliyor.',
    'Diğer onay modülleri (Envanter, Ödeme, Bordro Dönem, Avans) zaten muafiyetsizdi.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Normal onaycı — sıra kendisinde: Onayla/Reddet butonları görünür, işlem başarılı, akış sonraki adıma geçer.',
    'Kurucu yönetici — sıra başkasında: Listede talebi görür, butonlar çıkmaz. API ile onay denenirse “Bu onay adımı sizin pozisyonunuza ait değil” hatası; talep hiç değişmez.',
    'Kurucu yönetici — sıra kendisinde: Pozisyonuna atanmış adımı normal onaylayıp reddedebilmeli.',
    'Listeleme muafiyeti: Kendi pozisyonuna ait olmayan talepleri görmeye devam eder; istatistik ve dışa aktarma çalışır.',
    'Pozisyonu olmayan kullanıcı: “Employee position not found” hatası; kayıt değişmez.',
    'Sırası gelmemiş adım: Sonraki adımın onaycısı sırası gelmeden onaylayamaz.',
    'Zaten işlenmiş adım: Aynı adım ikinci kez onaylanamaz.',
])

# 2
add_section_heading(doc, '2. Ödeme Talebi — Fatura Görseli 50 MB + Alan Temizleme', page_break=True)
add_meta(doc, 'Ekran: Bordrom (Kişisel) → Ödeme Talebi Oluştur (Otomatik mod)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Görsel boyut sınırı 10 MB → 50 MB (hem okuma hem kaydetme).',
    'Okunamayan alanlar (tutar, KDV, tarih, açıklama) artık boşaltılıyor.',
    'Satıcı adı okuması sıkılaştırıldı: kısa işaretler tam kelime, ad doğru satırdan.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    '10–50 MB arası fotoğraf: Seçilir, okunur ve talep kaydedilir.',
    '50 MB üstü fotoğraf: “Görsel 50 MB sınırını aşıyor” uyarısı; dosya seçilmez.',
    'Alan temizleme: İlk fatura sonrası aynı pencerede tarihi okunamayan ikinci fatura → tarih boşalır.',
    'Hiçbir şey okunamayan görsel: Dört alan boşalır, “Faturadan bilgi okunamadı” uyarısı.',
    'Elle giriş modu: Görsel seçmek hiçbir alanı değiştirmez.',
    'Açıklama doğruluğu: Ya doğru firma adı ya boş; yanlış metin olmamalı.',
    'Sunucuda OCR yoksa: Uyarı çıkar, kullanıcı elle doldurur.',
])

# 3
add_section_heading(doc, '3. Dosya Yükleme — Sisteme 50 MB, Depolamaya 5 MB, Otomatik Küçültme', page_break=True)
add_meta(doc, 'Ekran: Dosya yüklenen tüm ekranlar (Doküman, Çalışan kartı, İzin ekleri, Ödeme, Avans, Duyuru, Bordro arşivi, CV, Envanter, profil fotoğrafı vb.)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Sisteme kabul sınırı 50 MB (eskisi 10 MB).',
    '5 MB üstü dosyalar otomatik küçültülüyor; görsellerde biçim korunuyor, PDF yeniden yazılıyor.',
    'Word/Excel/zip küçültülemez — olduğu gibi kaydedilir (50 MB’a kadar).',
    'Kullanıcıya “X dosya küçültüldü (42 MB → 2.8 MB)” bildirimi çıkar.',
    'Küçültme tarayıcıda ve sunucuda çift katmanlı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Büyük fotoğraf (30–45 MB): Küçültme bildirimi çıkar, yükleme başarılı, dosya okunabilir/net.',
    '50 MB üstü dosya: Reddedilir, “50 MB sınırını aşıyor” uyarısı.',
    'Biçim korunumu: PNG → PNG, JPG → JPG; şeffaf PNG’de şeffaflık bozulmaz.',
    'Büyük PDF (20 MB, taranmış): Açılır, sayfa sayısı aynı; metin PDF’de yazı seçilebilir.',
    'Word/Excel: Değişmeden kaydedilir — indirilen dosya birebir aynı.',
    'Küçük dosyalar (5 MB altı): İşlem yok, bildirim yok, dosya birebir korunur.',
    'Toplu yükleme: Çok dosya aynı anda yüklenebilir; sayı sınırı uyarısı çıkmaz.',
    'Profil fotoğrafı: Büyük fotoğrafta kırpma ve kayıt çalışır.',
    'İndirme/görüntüleme regresyonu: Küçültülen dosyalar Görüntüle/İndir ile açılır; zip bozulmaz.',
])

# 4
add_section_heading(doc, '4. Ödeme Talebi — Fatura OCR İyileştirmeleri', page_break=True)
add_meta(doc, 'Ekran: Bordrom → Ödeme Talebi Oluştur → Otomatik mod')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Açıklamadaki tek harfli gürültü temizleniyor.',
    'Ticari ad, hukuki unvan yerine doğru satırdan okunuyor.',
    'Kalem tutarının KDV sanılması önlendi.',
    'Tarihte boşluk toleransı (“19/ 08/ 2026”).',
    'En çok alan okuyan kip seçiliyor (erken durma hatası giderildi).',
    '“TOPLAM” etiketi öncelikli; tarih POS damgasından da okunabiliyor.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Yazarkasa fişi (tek KDV oranı): Tutar, KDV, tarih, açıklama dolu; tek harfli artık yok.',
    'Ticari adı ayrı satırda: Açıklamada adres/unvan değil ticari ad.',
    'Karma KDV oranlı fiş: KDV okunamazsa boş kalır, yanlış tutar yazılmaz.',
    'KDV satırı olmayan fiş: Tek oran varsa toplamdan hesaplanır; çok oran varsa boş.',
    'Eğik/kıvrık fotoğraf: Okunamayan alanlar boş; yanlış değer yok.',
    'Regresyon: Daha önce doğru okunan faturalarda sonuç değişmez.',
])

# 5
add_section_heading(doc, '5. Fatura Görseli — Modal İçi Sağ Panel + Büyüteç', page_break=True)
add_meta(doc, 'Ekran: Bordrom → Ödeme Talebi Oluştur / Ödeme Talebi Detayları')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Fatura görseli modal içinde sağ panelde (tam ekran katman kaldırıldı).',
    'Sol: talep bilgileri/form · Sağ: fatura · Arada dikey ayraç.',
    'Tıklayınca imleci takip eden kare büyüteç; tekrar tıklayınca kapanır.',
    '“Faturayı Görüntüle” butonu ve 96×96 küçük görsel kaldırıldı.',
    'Dar ekranda paneller alt alta.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Detay modalı (faturalı): İki panelli geniş modal; “Faturayı Görüntüle” yok.',
    'Detay modalı (faturasız): Eski dar tek panel; sağ panel/ayraç yok.',
    'Oluşturma modalı: Fatura seçilmeden dar; seçilince sağ panel açılır.',
    'Büyüteç: Bir tık aç, tekrar tık kapat.',
    'Hızlı çift tıklama: Kare açılır ve açık kalır.',
    'İmleç görselden çıkınca büyüteç kaybolur; geri girince görünür.',
    'Uzun sol içerik: Sol kaydırılırken sağ görsel yerinde kalır.',
    'Dar ekran: Paneller alt alta; yatay kaydırma yok.',
    'Dokunmatik: Dokununca açılır, parmak takip eder, sayfa kaymaz.',
    'Tema: Karanlık/aydınlıkta ayraç, çerçeve, büyüteç okunur.',
    'Uzun/dar fatura: Dikey fiş tamamı görünür; geniş boşluk kalmaz.',
    'Ekleri İndir butonu eskisi gibi çalışır.',
])

# 6
add_section_heading(doc, '6. Boşalan Pozisyon Onay Akışları — İki Sekmeli Toplu Düzenleme', page_break=True)
add_meta(doc, 'Ekran: Çalışan kartı → İşten Çıkar · Organizasyon → Pozisyonlar → çalışan değiştirme')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'İki sekme: “Çözüm zorunlu” ve “İsteğe bağlı”.',
    'Akışın tüm zinciri düzenlenebilir (taşı, değiştir, çıkar, ekle).',
    'Değişiklikler birikir; “Uygula” ile 20’şer grup halinde gönderilir.',
    'Toplu işlem tikleri sekmeye özel; arama kutusu; son adım koruması.',
    'Karar bekleyen akış varken “Uygula” pasif.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'İki sekme ayrımı: Tek adımlı üst akışsız → Çözüm zorunlu; yanında başka onaylayıcı → İsteğe bağlı.',
    'Aynı kişi akışta birden çok kez: Listede bir kez, Çözüm zorunlu; tüm adımlar kırmızı “Boşalıyor”.',
    'Karışık akış (Burcu→Arda→Arda, Arda boşalıyor): İsteğe bağlı; uygulanınca yalnız Burcu kalır.',
    'Bekleyen talep / rozet sayıları katlanmaz.',
    'Sekme içeriği karışmaz: Her sekme kendi akışlarını gösterir.',
    'Hiç dokunmadan uygulama: Yalnız isteğe bağlı varsa Uygula aktif; boşalan adım düşer.',
    'Karar beklerken Uygula pasif; tooltip kalan sayıyı yazar.',
    'Son adım çıkarılamaz: Tek adımlı akışta çıkar kapalı, tooltip neden yazar.',
    'Devralma bilgisi: Üst akışlı tek adım çıkarılınca zincir bilgisi görünür.',
    'Sıralama: Adım yukarı taşı → numaralar yenilenir; ayarlarda aynı sıra.',
    'Ekleme: İki onaylayıcı ekle + boşalanı çıkar → ayarlarda doğru yansır.',
    'Aynı pozisyon zincire iki kez eklenemez.',
    'Geri al: Satır ilk hâline döner; “Düzenlendi” kaybolur.',
    'Toplu işlem: Sekme değişince seçim temizlenir; diğer sekmeye uygulanmaz.',
    'Toplu çıkarmada engel: Çıkarılamayanlar atlanır; uyarı çıkar.',
    'Arama: Birim parçası süzülür; “Tümünü seç” yalnız süzülmüşleri seçer.',
    '100+ akış: İlerleme göstergesi çalışır; zaman aşımı olmaz.',
    'Kısmi başarısızlık: Reddedilen listede kalır; diğerleri uygulanır; uyarı çıkar.',
    'Bekleyen talepler: “Yeniden değerlendirmeye alındı” listesi + bildirim.',
    'Vazgeçme (İptal): Hiçbir akış değişmez; işten çıkarma yapılmaz.',
    'API ile penceresiz işten çıkarma → 409.',
    'Tema: Kırmızı zemin ve rozetler okunur.',
    'Dil: Altı dilde (tr/en/es/fi/ar/am) metinler çevrilmiş; İngilizce kalıntı yok.',
])

# 7
add_section_heading(doc, '7. Toplu Dosya Yükleme — Çalışan Seçimi Ortak Alana Geçti', page_break=True)
add_meta(doc, 'Ekran: Doküman Yönetimi → Ekip Dosyaları → Toplu İşlemler → Dosya Yükleme')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Ortak çalışan seçimi: profil fotoğrafı + pozisyon adı.',
    'Arama ad soyad ve pozisyon üzerinden.',
    'Sicil numarasıyla arama kalktı (bilinçli).',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Liste: Avatar + ad soyad + altında pozisyon; üst üste binme yok.',
    'Seçim: Kapalı kutuda “Ad Soyad - Pozisyon” tek satır; taşma yok.',
    'Ad/pozisyon araması çalışır (büyük/küçük harf fark etmez).',
    'Sicil numarası yazınca sonuç çıkmaz (beklenen davranış).',
    'Pozisyonsuz çalışan: Yalnız ad soyad; boş “ - ” artığı yok.',
    'Fotoğrafsız: Kırık görsel değil, ikonlu avatar.',
    'Çalışan değiştirince kartlar yenilenir; önceki dosyalar taşınmaz.',
    'Yetkisiz kullanıcıda liste boş; hata çıkmaz.',
    'Detay sayfasından açılınca alan kilitli; çalışan değiştirilemez.',
    'Karanlık tema: Avatar, ad, pozisyon okunur.',
])

# 8
add_section_heading(doc, '8. Şifreli Ad/Soyad İki Yerde Çözülmeden Görünüyordu', page_break=True)
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Anket atama mailinde selamlamadaki ad soyad düzgün (base64 düzeltildi).',
    'Arka plan iş göstergesinde de ad düzeldi.',
    'Bordro dönem çakışma uyarısında isimler okunabilir.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Anket atama maili: Ad soyad okunabilir; base64 metin yok.',
    'İş göstergesi: Gönderim sırasında çalışan adı okunabilir.',
    'Boş ad/soyad: Mailde fazladan boşluk/tire düşmez.',
    'E-posta tanımsız: Mail atlanır; iş hata vermeden biter.',
    'Pasif/işten ayrılmış: Mail gitmez.',
    'Türkçe karakter (ç ğ ı ö ş ü): Mailde bozulmaz.',
    'Bordro dönem çakışması: Uyarıda ad soyad ve birim okunabilir; kayıt engellenir.',
    'Mevcut döneme çalışan ekleme: Aynı uyarı, aynı okunabilirlik.',
    'Çakışma yoksa uyarı çıkmaz (regresyon).',
])

# 9
add_section_heading(doc, '9. İndir → İncele Butonu — Dosya İnceleme Penceresi', page_break=True)
add_meta(doc, 'Ekran: Dosya indirme butonu olan tüm ekranlar (Doküman, Çalışan, İzin, Ödeme, Avans, Duyuru, Bordro, CV, Envanter, Geri Bildirim, Ayarlar vb.)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Tüm indirme butonları “İncele” oldu (turuncu, büyüteçli ikon).',
    'Tek resim → pencere · Tek PDF/txt/xml/json → yeni sekme · Tek Word/Excel → indir.',
    'Çok dosya → pencere, oklar, sayaç, Tümünü İndir (zip).',
    '“İndirmek istediğinize emin misiniz?” onayları kaldırıldı.',
    '“Görüntüle + İndir” ikilileri tek İncele’de birleşti.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Tek resim: Pencere açılır; Kapat/İndir/Aç var; “Tümünü İndir” yok.',
    'Tek PDF: Pencere açılmadan yeni sekmede açılır.',
    'Açılır pencere engeli: Pencere açılır; “Aç” ile PDF yeni sekmede açılır.',
    'Tek Word/Excel: Pencere açılmadan iner; doğru ad/uzantı; zip değil.',
    'Çok dosyalı: Oklar, sayaç 1/3→2/3→3/3; sonda başa döner.',
    'Resimde Aç+İndir görünür; Word’deyken Aç görünmez.',
    'Tümünü İndir: Tek zip; gerçek adlar; aynı adlılar “ad (2).uzantı”.',
    'Tek dosyalı klasör: Tekli davranış (zip inmez).',
    'Boş/silinmiş dosya: “Görüntülenecek dosya bulunamadı”; sayfa donmaz.',
    'Yetkisiz: İncele görünmez; API reddeder.',
    'Onay penceresi kalktı: Ekip Dosyaları/Talepler/Gönderimler’de çıkmaz.',
    'Görüntüle+İndir birleşmesi: Uyarılar, Arşiv Bordrolar, Bordrolarım, avans dekontlarında tek buton.',
    'Avans dekontu: Kişisel ve yönetimden incelenir; başkasınınkine erişilemez.',
    'Eski ödeme belgeleri açılır/iner.',
    'Toplu İşlemler → Dosya İndirme: Satır İncele tekil; üst Tümünü İndir değişmemiş.',
    'Mobil: İncele İşlemler sütununda ikon; pencere yatay kaymaz.',
    'Karanlık tema: Pencere, dosya adı, sayaç, butonlar okunur.',
    'Dil: Altı dilde İncele ve pencere metinleri çevrilmiş.',
    'Bellek: İleri-geri gezip kapatıp başka kayıtta açınca yavaşlama yok.',
])

# 10
add_section_heading(doc, '10. Havuz (Ortak Kullanım) Talebi — Uçtan Uca Yeni Akış', page_break=True)
add_meta(doc, 'Ekranlar: Stok → Kategoriler/Ürünler · Zimmetlerim → Havuzdan Talep Et / Taleplerim · Zimmet Yönetimi → Talepler · Stok Talepleri')
add_meta(doc, 'Migration: 20260902112722_PoolInventoryAdded (uygulanmadan ekranlar açılmaz)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Tarih–saat aralıklı ortak kullanım ürünü talebi (araç, oda vb.).',
    'Akış zimmet ile aynı: Zimmet İş Akışı → Stok İş Akışı.',
    'Kategori: Havuza Açık + izinli birimler/çalışanlar; ürün: Havuz Ürünü.',
    'Doluluk takvimi; dolu araç seçilemez; çakışma sunucuda.',
    'Birleşik tablolar (Tür kolonu + süzgeç); otomatik iade (15 dk cron).',
])
add_h3(doc, 'Kategori / Ürün Ayarı')
add_check_group(doc, [
    'Havuz kapalı kategoride Havuz Ürünü alanı gri; API ile gönderilirse reddedilir.',
    'Havuz aç → ürün işaretle → kapat: Kapalıyken o kategoride talep açılamaz.',
    'İzin listesine üst birim: Alt birim çalışanı talep açabilir.',
    'İzin listesine tek çalışan: Aynı birimdeki başka çalışan açamaz / görmez.',
    'İzin listesini boşalt: Herkes tekrar açabilir.',
    'Havuzu kapatınca izin listesi temizlenir; tekrar açınca boş gelir.',
])
add_h3(doc, 'Talep Oluşturma')
add_check_group(doc, [
    'Havuz kategorisi yokken “Havuzdan Talep Et” butonu görünmez.',
    'Bitiş < başlangıç → hata; talep oluşmaz.',
    'Başlangıç = bitiş → hata.',
    'Ek sürücü = Evet, çalışan seçilmeden → hata.',
    'Başka kategorinin ürünü (API) → reddedilir.',
    'Tercih boş bırakılarak talep oluşturulabilir.',
    'Taleplerim’de Tür = Havuz; miktar yerine tarih aralığı.',
])
add_h3(doc, 'Onay Akışı')
add_check_group(doc, [
    'Zimmet İş Akışı tanımlı: Önce yönetici; onaydan sonra stok isteği.',
    'Zimmet İş Akışı tanımsız: Doğrudan stok isteği.',
    'Stok İş Akışı tanımsız: Talep oluşmaz; anlaşılır hata.',
    'Sırası gelmemiş adım onaylanamaz.',
    'Pozisyonu olmayan kullanıcı API’den onaylayamaz.',
    'Yönetici reddi: Reddedildi; takvimde doluluk yok.',
    'Stok görevlisi üstlendikten sonra çalışan iptal edemez.',
])
add_h3(doc, 'Stok Tarafı ve Teslim')
add_check_group(doc, [
    'Teslim listesi yalnız havuz ürünlerini gösterir.',
    'Tercih edilen araç formda hazır gelir; değiştirilebilir.',
    'Tercih edilen araç bakımda → alan boş açılır; hata yok.',
    'Çakışma: Aynı araç aynı saat → reddedilir; farklı gün → çalışır.',
    'Sınır (12:00 bitiş / 12:00 başlangıç): Çakışma sayılmaz.',
    'Stok red → araç Müsait’e döner.',
    'Stok görevlisi iptal → istek Beklemede; başka görevli üstlenebilir.',
])
add_h3(doc, 'Takvim')
add_check_group(doc, [
    'Yönetim takviminde dolu hücre: Kim + hangi saatler.',
    'Çalışan takviminde: Yalnız “Dolu”; isim/amaç/güzergâh yok.',
    'Yetkisiz kategori ürünleri çalışan takviminde yok.',
    'Bakımdaki ürün tüm günlerde “Bakımda”.',
    'Tercih edilmiş ama atanmamış talep takvimde görünmez; araç müsait.',
    'Reddedilmiş/iptal talepler takvimde yok.',
    'Gün / Hafta / Ay ve İleri / Geri / Bugün doğru aralık getirir.',
    'Departman süzgeci talep sahibinin birimine göre süzülür.',
])
add_h3(doc, 'Birleşik Tablolar')
add_check_group(doc, [
    'Tür = Havuz: Yalnız havuz satırları; sayfa 1’e döner.',
    'Tümü + sayfalama: Tekrarlama/atlama yok; sıralama doğru.',
    'Arama + departman + durum + Tür birlikte çalışır.',
    'Toplu İşlem: İki tür bekleyenler; onay/red doğru uca gider.',
    'Stok Talepleri kartları iki tür toplamını gösterir.',
])
add_h3(doc, 'Teslim / İade Otomatiği')
add_check_group(doc, [
    'Bitiş geçmiş + Zimmetli → en geç 15 dk içinde İade Bekleniyor.',
    'Bitiş geçmiş + hiç teslim alınmamış → tahsis iptal; ürün Müsait.',
    'Aynı kayıt ikinci turda tekrar işlenmez.',
    'Sunucu yeniden başlatılınca geçmiş kayıtlar sonraki turda yakalanır.',
])
add_h3(doc, 'Mesaj ve Yorum')
add_check_group(doc, [
    'Havuz yazışması çalışır; okunmamış rozeti doğru; modalda sıfırlanır.',
    'Zimmet mesajları havuzda (ve tersi) görünmez.',
    '“Talep sahibi görsün” işaretli değilse çalışan yorumu görmez.',
])

# 11
add_section_heading(doc, '11. Aynı Gün Düzeltmeleri (02 Eylül Test Geri Bildirimi)', page_break=True)
add_subheading(doc, '11.1 Havuz Ürünü Tiki Düzenlemede Boş Geliyordu')
add_check_group(doc, [
    'Havuz ürünü tikini aç → kaydet → tekrar düzenle: Tik açık gelir.',
    'Ürün Detayları’nda “Havuz Ürünü: Havuz” görünür.',
    'Kategori havuza kapalıyla değiştirilince bayrak sessizce sıfırlanır.',
])
add_subheading(doc, '11.2 Kategoriler Artık Birim Bazlı')
add_check_group(doc, [
    'Kök birimde kategori var, alt birimde yok → alt birimde Kategoriler boş.',
    'Alt birimde ürün eklerken kategori listesi boş.',
    'Kategori detayı ürün sayısı = Ürünler sekmesi sayısı.',
    'Silinmiş ürün kategori detayında görünmez.',
])
add_subheading(doc, '11.3 “Bilinmeyen” Onay Adımı Düzeltildi')
add_check_group(doc, [
    'İki adımlı yeni talep: 1. Beklemede, 2. Sırası Gelmedi.',
    '1. adım onaylanınca 2. adım Beklemede olur.',
    'Talep iptal → adımlar İptal Edildi.',
])
add_subheading(doc, '11.4 Tablolarda Talep Tarihi, Detayda İki Tarih')
add_check_group(doc, [
    'Talep açıp akışı değiştirince tabloda Talep Tarihi değişmez.',
    'Talepler sekmesinden detay: İki tarih; Oluşturulma = akış değişikliği tarihi.',
    'Taleplerim’den detay: Yalnız Talep Tarihi.',
    'Tarih kolonu sıralaması talep tarihine göre.',
    'Dışa aktarmada Talep Tarihi doğru.',
])
add_subheading(doc, '11.5 Akış Değişince Havuz Onayları da Yeniden Kuruluyor')
add_check_group(doc, [
    'Bekleyen havuz + Zimmet İş Akışı değişikliği: Uyarıda “Havuz Talebi” etiketi.',
    'Kaydet → havuz onay zinciri yeni akışa göre kurulur.',
    'Zimmet talebi regresyonu: Eskisi gibi çalışır.',
])

# 12
add_section_heading(doc, '12. Havuz Takvimi — Yönetim/Çalışan Ayrımı ve Çakışma Uyarısı', page_break=True)
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Yönetimde onay bekleyenler sarı; teslim edilenler kırmızı.',
    'Aynı hücrede çok kayıt: sayaç + “+N” + detay paneli.',
    'Çakışma bilgilendirme şeridi (yalnız uyarı; kilitlemez).',
    'Yeni: Ürün Yönetimi → Havuz Takvimi sekmesi.',
    'Talep satırından açılan takvimde kategori kilitli.',
    'Çalışanda “Onay Bekliyor” hiç yok; durum tahsis durumundan.',
    'Kategori süzgeci seçili birime göre; ürün birimi ≠ talep departmanı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Onay bekleyen — çalışan: Günler Müsait; renk açıklamasında Onay Bekliyor yok.',
    'Onay bekleyen — yönetim: Günler sarı; “(onay bekliyor)” yazısı.',
    'Onayla + teslim et → hücreler kırmızı.',
    'Çalışan teslim hücre tıklama: “Rezerve”; kişi/amaç/güzergâh yok.',
    'Aynı araca 3 onay bekleyen: “3 talep”, 2 liste + “+1”; panelde 3 kayıt.',
    'Çakışma şeridi: Çakışmada çıkar; yoksa çıkmaz.',
    'Kategori: Talep satırından kilitli; Ürün Yönetimi’nden serbest.',
    'Birim değiştir → kategori listesi ve ürün satırları yenilenir.',
    'Departman süzgeci ürün değil talep sahibine göre süzülür.',
    'Sınır 17:00 bitiş/başlangıç: 2 kayıt görünür; çakışma şeridi yok.',
    'Bakımdaki araç: Gri, tıklanamaz; onay bekleyen olsa bile sarıya dönmez.',
    'Kategori seçmeden: Seçili birimin tüm havuz ürünleri listelenir.',
    'Talep reddi → hücre yeşil; kayıt düşer.',
    'Onay bekleyen iptal → hücre yeşil.',
    '62 günden uzun aralık → 62 günle sınırlanır; donma yok.',
    'Koyu tema: Sarı/kırmızı/yeşil/gri hücreler okunur.',
])

# ═══════════════════════════════════════════
# 13 — 2026-09-03
# ═══════════════════════════════════════════
add_section_heading(doc, '13. “İncele” Her Dosya Türünde Önce Önizleme Penceresini Açıyor', page_break=True)
add_meta(doc, 'Ekran: İncele butonu olan tüm listeler · özellikle Toplu İşlemler → Dosya İndirme · Özlük Bilgileri · Mutabakat')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Dosya sayısı/türü ne olursa olsun İncele her zaman önce önizleme penceresini açar.',
    'Word/Excel artık kendiliğinden inmez; PDF kendiliğinden yeni sekmeye gitmez — karar pencereden verilir.',
    'Resim: görsel + İndir/Aç/Kapat · PDF/txt/xml/json: “yeni sekmede açılır” + İndir/Aç/Kapat · Diğer (xlsx/docx/zip/uzantısız): “görüntülenemiyor” + yalnız İndir/Kapat (Aç yok).',
    'Çoklu dosyada oklar, sayaç ve Tümünü İndir (zip) eskisi gibi.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Dosya İndirme → tek Excel: Kendiliğinden inmez; pencere açılır; “görüntülenemiyor”; Aç yok; İndir doğru adla iner.',
    'Tek PDF: Yeni sekme kendiliğinden açılmaz; pencerede “yeni sekmede açılır”; Aç ile PDF açılır; pencere kapanınca sekmedeki PDF bozulmaz.',
    'Tek resim: Pencerede görsel; Aç ve İndir çalışır.',
    'Uzantısız dosya: Aç yok, İndir çalışır; hata/boş pencere yok.',
    'Çoklu dosya: Oklar, sayaç, Tümünü İndir zip; aynı adlılar “ad (2).uzantı”.',
    'Boş/silinmiş dosya: “dosya bulunamadı”; boş pencere açılmaz.',
    'Aynı satıra arka arkaya: Kapatıp tekrar açınca dosya yine görünür.',
    'Özlük Bilgileri: Aynı davranış Dosya İndirme ile birebir.',
])

# 14
add_section_heading(doc, '14. Onay Akışı Bildirimlerinde “poolRequest” → “Havuz Talebi”', page_break=True)
add_meta(doc, 'Ekran: Duyurular / Duyurularım · e-posta · tetikleyen: Pozisyon boşaltma / onay akışı değiştirme')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Havuz talepleri bildirimde ham “poolRequest” yerine “Havuz Talebi” / “Pool Request” yazıyor.',
    'Diğer talep türleri zaten çevriliydi; sorun yalnız duyuru ve e-postadaydı.',
    'Eski duyurular DB’de metin olarak kaldığı için hâlâ poolRequest gösterebilir.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Havuz talebi + akış değişikliği: Duyuruda “[…] Havuz Talebi — …”; poolRequest yok.',
    'Aynı senaryoda e-posta Türkçe satır.',
    'İngilizce tenant: “[Inventory] Pool Request — …” (veya güncel Zimmet etiketiyle).',
    'Zimmet + havuz birlikte: İki satır da çevrili; birbirine karışmaz.',
    'Yeni onaylayıcı bildiriminde tür adı çevrili.',
    'Çok talepli bildirimde listelenen satırlarda ham kod adı yok.',
])

# 15
add_section_heading(doc, '15. Satır Hover Grileşmesi Düzeltildi (5 Ekran)', page_break=True)
add_meta(doc, 'Ekran: Ekip Dosyaları · Avans (5 sekme) · Ek Mesai → Talepler · Vardiya → Talepler · Vardiya → Uyum → İhlaller')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Hover’da artık tüm satır grileşiyor (önceden yalnız İşlemler kolonu).',
    'Başlık satırı gri zemine kavuştu; karanlık temada da geçerli.',
    'Punto/hücre boşluğu sistemle eşitlendi; sabit İşlemler kolonu davranışı aynı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Ekip Dosyaları: Satırın tamamı grileşir.',
    'İşlemler kolonu satırın geri kalanından bir tık daha koyu (bilinçli).',
    'Başlık satırı gri zeminli.',
    'Karanlık tema: Hover ve başlık koyu; beyaz/şeffaf leke yok.',
    'Renkli satırlar (Ek Mesai/Vardiya Talepler): Hover’da grileşir (renk kaybolabilir — beklenen).',
    'Yatay kaydırma: İşlemler kolonu arkası opak.',
    'Mobil: İşlemler butonları kırpılmaz; sütun ekranın yarısını kaplamaz.',
    'Ek Mesai → Talepler “Süre”: İkon + metin tek satırda yan yana.',
    'Avans Yönetimi: Beş sekmenin hepsinde hover düzgün.',
])

# 16
add_section_heading(doc, '16. Tablo Yoğunluğu Tüm Sistemde Eşitlendi (68 Tablo)', page_break=True)
add_meta(doc, 'Ekran: Duyurular, CV Analiz, Doküman, Bilgilerim İzinler, Envanter kişisel, Bordro, Performans, Raporlar, Vardiya, Stok, Sistem Yönetimi, Ayarlar altındaki birçok tablo')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    '68 tablo “sıkışık” yoğunluğa alındı; satır yüksekliği ve hücre boşluğu azaldı.',
    'İçerik, kolonlar, sıralama, filtreler, butonlar değişmedi.',
    'Bütçe, Bordro/SGK beyanname, onay akışı tanım ve mobilde özel ayarlı tablolara dokunulmadı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Listedeki tablolar İzin Yönetimi → Talepler ile aynı satır yüksekliğinde.',
    'Uzun içerik kırpılmaz / üst üste binmez.',
    'İşlem ikonları dikeyde ortalı; satırı taşırmıyor.',
    'Genişletilebilir satırlar (Envanter Paketleri, Performans Akışları) hizalı.',
    'Sayfalama ve sıralama veri/sıra bozulmadan çalışır.',
    'Mobil: Tablo sayfayı yana itmez; İşlemler sağda kalır.',
    'Karanlık tema: Zemin ve ayraç diğer tablolarla aynı.',
])

# 17
add_section_heading(doc, '17. Bildirimlerde “Envanter” → “Zimmet”', page_break=True)
add_meta(doc, 'Ekran: Zil / mobil push · Duyurular / Duyurularım · e-posta (zimmet talebi oluşturma/onay/red · akış değişikliği)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Bildirim metinleri arayüzle aynı: “Zimmet” / “Assignment” (eski: Envanter / Inventory).',
    'Yeni talep, onay/ret, akış değişikliği duyurusu (modül + talep türü) güncellendi.',
    'Diğer modül bildirimleri değişmedi. Eski kayıtlar metin olarak eski kelimeyi tutabilir.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Yeni zimmet talebi: Zil başlığı “Yeni Zimmet talebi”; Envanter yok.',
    'Onay/red: “Zimmet talebiniz onaylandı/reddedildi”.',
    'Push ile zil metni aynı.',
    'Akış değişikliği: “[Zimmet] Zimmet Talebi — …” ve havuz için “[Zimmet] Havuz Talebi — …”.',
    'İngilizce: Assignment / Assignment Request; Inventory yok.',
    'İzin, ek mesai, avans, döküman, performans başlıkları birebir aynı (regresyon).',
])

# 18
add_section_heading(doc, '18. Stok Taleplerinde “Teslim Aşamasında” Ara Durumu', page_break=True)
add_meta(doc, 'Ekran: Stok Talepleri → Talep Onayları (Zimmet + Havuz) · Taleplerim detay zaman çizelgesi')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'İşlemler → Teslim Et artık doğrudan “Teslim Edildi” yapmıyor → “Teslim Aşamasında” (turuncu).',
    'Yeşil tik ile ürün tanımlanır; Tamamla → Teslim Edildi + Talep Onaylandı.',
    'Eski “Teslim Edildi ama tamamlanmamış” satırlarda yeşil tik/Tamamla çalışmaya devam eder.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Zimmet: Teslim Et → Durum Teslim Aşamasında; Talep Durumu hâlâ Beklemede.',
    'Ürün ekle, pencereyi kapat: Hâlâ Teslim Aşamasında (Teslim Edildi olmamalı).',
    'Tamamla → Teslim Edildi + Onaylandı; yalnız Görüntüle kalır.',
    'Havuz: Aynı akış; tüm araçlar teslim edilmeden Tamamla pasif + uyarı.',
    'Yarım bırakma: Teslimi İptal Et → İşlemde; İşlemler butonu yeniden çıkar.',
    'Ürün varken teslim iptali (API) reddedilmeli.',
    'Tamamlandıktan sonra durum değişmez; yenilemede Teslim Edildi kalır.',
    'Transfer: Yalnız teslim aşamasındaki satır Teslim Edildi olur; eski transfer satırları korunur.',
    'Teslim aşamasında Reddet → Reddedildi; asılı kalmaz.',
    'Taleplerim detayında rozet: Teslim Aşamasında → sonra Teslim Edildi.',
    'Altı dilde rozet çevrili; ham anahtar (inDelivery) yok.',
    'Eski kayıt: Yeşil tik ve Tamamla hâlâ çalışır.',
])

# 19
add_section_heading(doc, '19. Ürün Geçmişi — Saat Gösterimi ve Sıralama Düzeltmesi', page_break=True)
add_meta(doc, 'Ekran: Stok → Ürünler → Geçmiş çekmecesi (+ “i” detay)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Her olayda tarih + saat (GG.AA.YYYY SS:dd).',
    'Sıralama sisteme yazılma anına göre; “Ürün Kayıt Edildi” en altta.',
    'Atama/Geri Alma Tarihi detayda ayrı; Kayıt Tarihi detaya eklendi.',
    'Aynı saniyedeki olaylar tutarlı sırada.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Yeni ürün + ata: “Ürün Kayıt Edildi” en altta; üstte sonraki olaylar.',
    'Her satırda saat; Invalid Date yok.',
    'Aynı gün ata → geri al → ata: Saat sırasına göre yeniden eskiye.',
    'Geçmiş tarihli atama: Çizelgede işlem anı; detayda Atama Tarihi = seçilen gün.',
    'Kayıt Tarihi detayda görünür ve seçilen günü yazar.',
    'Aynı saniye olayları her açılışta aynı sırada.',
    'Tümü / Zimmet-İade / Ticket / Bakım / Satın Alma sekmelerinde sıra kuralı aynı.',
    '“Ürün Kayıt Edildi” hiçbir zaman başka olayın üstünde değil.',
    'Boş geçmiş: Tek satır + saat.',
    'Karanlık tema / mobil: Saatli tarih taşmaz.',
])

# 20
add_section_heading(doc, '20. Zorunlu Alan Uyarıları — Çift Bildirim, Çalışan Güncelle, Bilgilerim Opsiyonel', page_break=True)

add_subheading(doc, '20.1 Yeni Çalışan Ekle — Uyarı İki Kez Çıkıyordu')
add_meta(doc, 'Ekran: Personel Yönetimi → Yeni Çalışan Ekle → Kaydet')
add_check_group(doc, [
    'Zorunlu alan boş → Kaydet: Sağ üstte tek uyarı (çift değil).',
    'Üst üste 3 deneme: Her seferinde tek uyarı.',
    'Birden fazla sekmede eksik: Tek uyarı; rozetler/kırmızı çerçeveler; ilk eksik sekmeye geçiş.',
    'Hepsi doluyken: Çalışan bir kez oluşur; başarı bir kez; pencere kapanır.',
    'Olmaması gereken: Çift uyarı, çift başarı, çift kayıt.',
])

add_subheading(doc, '20.2 Çalışan Güncelle — Tüm Sekmelerde Zorunlu Alan Kontrolü')
add_meta(doc, 'Ekran: Personel Yönetimi → Çalışan Güncelle → Güncelle')
add_check_group(doc, [
    'Telefon sil → Güncelle: Tek uyarı, Temel Bilgiler “1” rozeti, kırmızı çerçeve.',
    'Telefonu doldur (Kaydet’siz): Rozet ve çerçeve anında kaybolur.',
    'Başka sekmedeyken zorunlu alan boş → Güncelle: O sekmeye otomatik geçiş.',
    'Farklı sekmelerde eksikler: Rozet sayıları doğru; ilk sekmeye geçiş.',
    'Dinamik (+) alan: Hiç adım yoksa eksik; en az bir adım doluysa değil.',
    'Koşula bağlı görünmeyen alan eksik sayılmaz.',
    'Dosya zorunlu: Belge varsa OK; silinince eksik.',
    'Hepsi doluyken Değişiklik Özeti açılır.',
    'Değişiklikleri Geri Al → rozet/çerçeve temizlenir.',
    'Olmaması gereken: Çift uyarı; eksikken özet penceresi; görünmeyen alan rozeti.',
    'Eski çalışan: Sonradan zorunlu yapılan boş alan güncellemeyi engeller.',
])

add_subheading(doc, '20.3 Bilgilerim — Hiçbir Alan Zorunlu Değil')
add_meta(doc, 'Ekran: Bilgilerim → Bilgilerimi Düzenle')
add_check_group(doc, [
    'Ayarlarda zorunlu alan boşken Bilgilerim’de Kaydet çalışır; uyarı yok.',
    'Aynı alan Ekle/Güncelle’de hâlâ zorunlu (rozet + çerçeve).',
    'Zorunlu dosya alanı belgesiz kaydedilebilir.',
    'Değiştirilemez alan kilitli/gri kalır.',
    'Olmaması gereken: “gereklidir” hatası; eksik alan yüzünden Kaydet engeli.',
])

# 21
add_section_heading(doc, '21. Yeni Çalışan Ekle — Hata Sonrası Mükerrer Kayıt Engeli', page_break=True)
add_meta(doc, 'Ekran: Personel Yönetimi → Yeni Çalışan Ekle → Kaydet')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Yan adım hatasında çalışan kaydı oluşmuşken tekrar Kaydet ikinci kopya yaratıyordu.',
    'Aynı temel bilgilerle tekrar denemede mevcut kayıt kullanılır; yalnız yarım adımlar tamamlanır.',
    'Mavi bilgi: “önceki denemede zaten oluşturulmuştu…”. Temel bilgi değişirse / Temizle / başarı sonrası hafıza sıfırlanır.',
    'Hata durumunda form temizlenmez. API benzersizlik kontrolü hâlâ yok (kapsam dışı).',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Yan adım hatası → değiştirmeden tekrar Kaydet: Mavi bilgi; listede tek çalışan.',
    'Aynı hatayı 3 kez: Yine tek çalışan.',
    'Hata sonrası Personel No (veya Ad/Soyad/E-posta/…) değiştir → yeni çalışan oluşur.',
    'Hata → Temizle → yeni kişi: Önceki kayda bağlanmaz.',
    'Başarılı kayıttan sonra yeni kişi: Öncekine bağlanmaz.',
    'Pozisyonda tek kişi; “ayrılmış” kopya yok.',
    'Olmaması gereken: Çoklu kopya; form verisinin hata sonrası silinmesi.',
    'Hatasız akış: Tek çalışan, yeşil başarı, pencere kapanır, liste yenilenir.',
])

# ═══════════════════════════════════════════
# 22–26 — 2026-09-04 Masraf
# ═══════════════════════════════════════════
add_section_heading(doc, '22. Masraf Modülü Genişletildi — Satıcı, Döviz, KDV, Etiket, Geri Ödeme', page_break=True)
add_meta(doc, 'Ekran: Bordrolarım → Ödeme Talepleri · Bordro Yönetimi → Talepler · Ayarlar → Masraf Kuralları')
add_meta(doc, 'Migration: AddExpenseManagement (zorunlu)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Yeni alanlar: Satıcı, para birimi/fiş tutarı/kur, KDV oranı, etiket, masraf merkezi, ödeme yöntemi, geri ödenecek, müşteriye yansıt, katılımcılar, masraf raporu.',
    '“Cebimden ödedim” kapalıysa onaylansa bile bordroya/puantaja düşmez (şirket kartı çift ödemeyi önler).',
    'PDF fatura yüklenebilir (OCR yok; sağ panel açılmaz).',
    'Listelerde Satıcı / Etiket / Geri Ödeme kolonları; “Ödeme Bekleyen” kartı; “Ödendi İşaretle”.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Ödeme tipi seç → Satıcı, Etiket, Masraf Merkezi, Ödeme Yöntemi görünür.',
    'Para birimi boş → Fiş Tutarı/Kur yok; Tutar elle yazılır.',
    'USD seç → alanlar çıkar, Tutar kilitli; 100 × 34,50 = 3.450,00.',
    'Kur boş gönderim → “Döviz seçtiyseniz kuru girin”.',
    'KDV %20, tutar 1.200 → KDV 200,00; elle 150 → sunucu reddeder (tolerans 1).',
    'KDV oranı boş, tutar 150 → kabul.',
    'KDV > toplam → reddedilir.',
    'Cebimden ödedim kapalı + onay → bordro/puantajda yok; listede var; “Geri ödeme yok”.',
    'Cebimden ödedim açık + onay → bordroya düşer; “Ödeme bekliyor”.',
    'Ödendi İşaretle → “Ödendi”; buton kaybolur; ikinci işaretleme hata.',
    'Geri ödeme yok satırda Ödendi İşaretle görünmez.',
    'PDF (otomatik): OCR yapılamaz uyarısı; sağ panel yok; elle gönderilebilir.',
    'JPEG (otomatik): tutar/KDV/oran/tarih/satıcı dolu; Açıklama boş.',
    'Katılımcı detayda etiket olarak görünür.',
    'Detayda satıcı, döviz, KDV oranı, etiket, merkez, yöntem, geri ödeme, katılımcılar.',
    'Tutar ≤ 0 → “sıfırdan büyük olmalı”.',
    'Para birimsiz kur (API) → “Kur yalnızca para birimiyle…” hatası.',
])

# 23
add_section_heading(doc, '23. Masraf Etiketleri (Muhasebe Kırılımı)', page_break=True)
add_meta(doc, 'Ekran: Ayarlar → Birim Yönetimi → Bordro Yönetimi → Masraf Etiketleri')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Ödeme tipinden bağımsız hiyerarşik etiket ağacı (Ad, Muhasebe Kodu, Üst, Sıra, Seçilebilir).',
    'Birim bazlı; tanımsız birimde üst birim etiketleri miras alınır.',
    'Alt etiketi olan veya masrafta kullanılan etiket silinemez; seçilebilir kapatılır.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Birim seç → boş liste + Yeni Etiket aktif.',
    'Birim seçmeden Yeni Etiket pasif.',
    'Üst + alt → listede “İdari Giderler:Kırtasiye”.',
    'Üstü sil → “Alt etiketleri önce silin”.',
    'Masrafta kullanılmışı sil → seçilemez işaretle uyarısı.',
    'Seçilebilir kapat → yeni formda yok; eski detayda görünür.',
    'Alt birimde etiket yok → üst birimin etiketleri.',
    'Alt birimde bir etiket tanımlı → yalnız o (miras kesilir).',
    'Etiket zorunlu kuralı açıkken etiketsiz masraf reddedilir.',
])

# 24
add_section_heading(doc, '24. Masraf Raporu — Birden Çok Masrafı Tek Pakette Onaya Gönderme', page_break=True)
add_meta(doc, 'Ekran: Bordrolarım → Masraf Raporlarım (yeni sekme)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Taslak rapor (MR-…) altında masraflar toplanır; Onaya Gönder ile tek seferde akış kurulur.',
    'Taslak satırlar onaycıda/bordroda görünmez. Yarım gönderim yok — hata olursa hepsi taslağa döner.',
    'Rapor durumu satırlardan türetilir; silme/çıkarma yalnız taslakta.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Yeni rapor: Taslak, 0 satır, Onaya Gönder pasif.',
    '3 masraf ekle: Satır 3, toplam doğru, durum Taslak.',
    'Taslak satırlar onaycı kuyruğunda yok.',
    'Bordro hesapla → taslak satırlar düşmez.',
    'Onaya Gönder → Beklemede; onaycıda 3 kayıt; rapor Onayda.',
    'Gönderilmişte Masraf Ekle/Düzenle yok; API “yalnız taslak…” hatası.',
    '3 satır onay → rapor Onaylandı.',
    '2 onay + 1 red → rapor Onaylandı.',
    'Hepsi red → Reddedildi.',
    '1 onay + 1 beklemede → Onayda kalır.',
    'Geri ödenecekler Ödendi → rapor Ödendi; yarısı → Onaylandı kalır.',
    'Taslak sil → rapor + satırlar kaybolur.',
    'Gönderilmiş sil → “yalnız taslak…” hatası.',
    'Boş rapor gönder → “gönderilecek masraf yok”.',
    'Pozisyonsuz çalışan gönder → hata; hiçbir satır Beklemede olmaz.',
    'Başkasının rapor id’si → 403.',
    'Dönem başlangıcı > bitiş → reddedilir.',
])

# 25
add_section_heading(doc, '25. Masraf Kuralları (Politika) Ödeme Tanımına Bağlandı', page_break=True)
add_meta(doc, 'Ekran: Ayarlar → Ek Kazançlar/Kesintiler → “Çalışan giderine izin ver” → Masraf Kuralları')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Opsiyonel kurallar: üst limit, fiş eşiği, en eski fiş günü, varsayılan KDV, gelecek tarih, mükerrer engeli, satıcı/etiket/katılımcı zorunlu, müşteriye yansıt izni.',
    'Formda mavi özet + tarih seçicide kural dışı günler tıklanamaz.',
    'Yalnız çalışan talebinde geçerli; yönetim ek kazanç / Excel ithalat takılmaz.',
    '“Çalışan giderine izin ver” kapanınca kurallar varsayılana döner.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Gider izni kapalı → Masraf Kuralları bölümü yok.',
    'Açınca tüm alanlar boş/kapalı.',
    'Üst limit 500: 600 formda + API’de reddedilir.',
    'Fiş eşiği 200: 150 fişsiz OK; 250 fişsiz hata.',
    'En eski 30 gün: Eski günler tıklanamaz; API 40 gün öncesi red.',
    'Gelecek tarih kapalı → yarın seçilemez; açık → seçilebilir.',
    'Varsayılan KDV %20 → formda hazır gelir, değiştirilebilir.',
    'Mükerrer engeli: Aynı gün/tutar/satıcı ikinci kez red.',
    'İptal sonrası aynı masraf tekrar → kabul.',
    'Tutar 1 kuruş değişince kabul.',
    'Satıcı zorunlu → satıcısız red.',
    'Katılımcı zorunlu → katılımcısız red.',
    'Müşteriye yansıt kapalı → kutu yok; API isBillable true → red.',
    'Gider iznini kapat-aç → kurallar boş/kapalı.',
    'Yönetim ek kazanç limit aşarsa reddedilmez.',
])

# 26
add_section_heading(doc, '26. Masraf Girişi İkinci Tur — OCR/Döviz, Canlı Kur, Katılımcı Seçimi', page_break=True)
add_meta(doc, 'Ekran: Yeni Ödeme Talebi Oluştur · Masraf Raporu → Masraf Ekle (aynı pencere)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'OCR Satıcı’yı doldurur; Açıklama boş kalır.',
    'Para birimi değişince tutar Fiş Tutarı’na taşınır; kur 1→TCMB; geri dönüş tersine.',
    'Döviz seçiliyken OCR tutarı Fiş Tutarı’na yazılır; Tutar = × kur.',
    'Kur TCMB’den (bütçe ile aynı servis); yenile ikonu; alınamazsa 1 + uyarı, form bloklanmaz.',
    'Katılımcılar: İş arkadaşları kartı + serbest metin. Birim para birimi Genel Ayarlar’dan.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'OCR: Satıcı dolu, Açıklama boş.',
    'TRY birimde Tutar 1000 → USD: Fiş Tutarı 1000, Kur 1→TCMB, Tutar = 1000×kur.',
    'TRY’ye geri dön: Fiş/Kur kaybolur, Tutar 1000.',
    'USD→EUR: Yalnız kur tazelenir; fiş tutarı korunur.',
    'Yenile ikonu kuru ve bilgi satırını günceller.',
    'Ağ kesik: Kur 1, “alınamadı” uyarısı; form gönderilebilir.',
    'Önce USD sonra OCR: Tutar Fiş Tutarı’nda; Tutar = ×kur; KDV çevrilmiş tutardan.',
    'Birim para birimi USD: Tutar (USD); USD seçmek dönüşüm açmaz; TRY açar.',
    'Alt birimde tanım yok → üst birimin para birimi.',
    'Çalışan görme yetkisi var: İş Arkadaşları kartı + çoklu seçim.',
    'Yetki yok: Kart yok; yalnız serbest metin; konsol hatası yok.',
    '2 arkadaş + 1 dış → detayda 3 etiket.',
    'Katılımcı zorunlu: Yalnız serbest metin OK; hiçbiri → hata.',
    'Para birimi ayarı geç gelirse USD→USD: Dönüşüm kapanır; Tutar fiş tutarına döner.',
])

# 27
add_section_heading(doc, '27. Yeni Modül: Seyahat Yönetimi', page_break=True)
add_meta(doc, 'Ekran: Modül Merkezi → Seyahat Yönetimi · Kişisel → Seyahatlerim · Ayarlar → Seyahat Yönetimi')
add_meta(doc, 'Migration: AddTravelModule (AddExpenseManagement’tan SONRA). Yetkiler: ViewTravelManagement / ViewEmployeeTravels / ViewMyTravels')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Görev seyahati talebi: tür, tarih, şehir, ulaşım, konaklama, vekil, avans, amaç.',
    'Avans tutarı doluysa son onayda iş avansı otomatik açılır; boşsa avans doğmaz.',
    'Ödenmiş avansı olan seyahat iptal/silinemez. Tarih çakışması engellenir.',
    'Türler birim bazlı + üst birim mirası. Akış yoksa talep açılmaz.',
    'Yönetim başkası adına açsa da çalışanın onay zincirine girer. Süper admin onay muafiyeti yok.',
    'Kapsam dışı: belge eki, harcırah/döviz, puantaj/izin bağ, masraf raporu toplu bağ, toplu onay, dışa aktar.',
])
add_h3(doc, 'Talep Açma')
add_check_group(doc, [
    'Akışsız birimde talep açılmaz; akış tanımlanınca açılır.',
    'Bitiş < başlangıç engellenir.',
    '10–14 Eylül onaylı varken 12–16 talep → çakışma; 15–18 geçer.',
    'En fazla 3 gün türünde 5 günlük aralık form + sunucu reddeder.',
    'Aynı gün gidip dön = 1 gün; “en fazla 1 gün” türünde geçer.',
    'Vekil olarak kendini seçmek engellenir.',
    '“Çalışan talep edebilir” kapalı tür çalışan formunda yok; yönetimde var.',
    'Avans izni kapalı → avans alanı kaybolur ve tutar temizlenir.',
])
add_h3(doc, 'Onay Akışı')
add_check_group(doc, [
    'İki adımlı: ilk onay sonrası hâlâ Onay Bekliyor; ikinciye e-posta.',
    'Sırası gelmemiş ikinci onaylayan kuyrukta görmez.',
    'Pozisyonu olmayan kullanıcı onay ucu → 403.',
    'Redde gerekçesiz onay engellenir.',
    'Red sonrası sonraki adımlar sıraya girmez; sahibe gerekçeli e-posta.',
])
add_h3(doc, 'Avans Bağı')
add_check_group(doc, [
    'Avans boş + son onay → Avans Yönetiminde yeni kayıt yok.',
    'Avans dolu + son onay → iş avansı aynı tutarda; “Ödeme Bekliyor”.',
    'Avans ödenince seyahat rozeti “Ödendi”.',
    'Ödenmiş avanslı seyahat iptali engellenir; ödenmemişken iptal geçer.',
    'Akış değişip son adım ikinci kez onaylanırsa ikinci avans açılmaz.',
    'İş avansı kapalı birimde avanslı onay → seyahat onaylı kalır, avans açılmaz.',
])
add_h3(doc, 'Düzenleme / İptal / Türler / Yetki / Dil')
add_check_group(doc, [
    'Onaysız talep düzenlenir; ilk onay sonrası düzenle yok; API 400.',
    'İptal edilmiş talep bekleyen kuyrukta yok; onay/red engellenir.',
    'Alt birimde miras listesi salt okunur + uyarı; kendi tür eklenince miras kesilir.',
    'Silinen tür yeni taleplerde yok; eski taleplerde ad kalır.',
    'Yalnız ViewMyTravels: Seyahatlerim var, yönetim menüde yok, uç 403.',
    'Yalnız ViewEmployeeTravels: iki sekme var, silme yok.',
    'Yetkisiz Ayarlar’da Seyahat sekmesi yok.',
    'Birim yetkisi dışı çalışan aramada çıkmaz.',
    'Altı dilde başlık/sekme/kolon/rozet/hata çevrili.',
])

# 28
add_section_heading(doc, '28. Seyahat ↔ Masraf Bağı', page_break=True)
add_meta(doc, 'Ekran: Ödeme Talebi Oluştur → Seyahat alanı · Seyahat detayında maliyet paneli · listelerde Masraf kolonu')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Masraf satırına isteğe bağlı Seyahat bağlanır (rapora değil satıra).',
    'Listede yalnız oturumdaki çalışanın ONAYLANMIŞ seyahatleri; yoksa alan görünmez.',
    'Panel: Toplam / Cepten / Şirket / Seyahat Avansı. Otomatik mahsup YOK.',
    'İptal/red toplamdan düşer; taslak satırlar toplama girer.',
    'Sunucu: seyahat var mı, talebi açan kişiye mi ait, onaylanmış mı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Onaylı seyahati olan çalışan: Seyahat alanı görünür; yalnız onaylılar.',
    'Onaylı seyahati yok: alan yok; masraf normal çalışır.',
    'Bekleyen seyahat listede yok.',
    'Bağlı masraf detay panelinde görünür; Toplam artar.',
    'Cebimden açık → Cepten; kapalı → Şirketin Ödediği; toplam ikisinin toplamı.',
    'Masraf reddi / iptali toplamdan düşer.',
    'Taslak rapordaki masraf toplama girer; durum Taslak.',
    '3 masraf → liste kolonunda toplam tutar.',
    'Bağlı masrafı olmayan satırda Masraf = “-”.',
    'Avans + masraf kartları dolu; fark hesaplanmaz; mahsup bilgi kutusu var.',
    'Başkasının seyahat id’si ile masraf → 403.',
    'Onaylanmamış seyahat id → 400.',
    'Seyahat yetkisi olmayan masraf formu: alan yok, konsol hatası yok.',
    'Masraf yetkisi olmayan onaylayıcıda panel gizlenir; modal çalışır.',
    'Masraf detay/listesinde seyahat etiketi (başlık veya şehir).',
    'Altı dilde yeni alan/panel/kolon çevrili.',
])

# 29
add_section_heading(doc, '29. Seyahat Sekmeleri Yeniden Düzenlendi + Detay Ekranı', page_break=True)
add_meta(doc, 'Ekran: Seyahat Yönetimi (Talepler / Ekip Seyahatleri / Seyahat Giderleri) · Seyahat Detayı')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Sayfa Talepler ile açılır; Ekip Seyahatleri eski Seyahatler; yeni Seyahat Giderleri.',
    'Giderler’de yalnız avansı veya masrafı olan seyahatler. Kartlar süzgecin tamamı (sayfa değil).',
    'Detay: Ad Soyad - Seyahat Detayı; Alan/Değer; Onay Akışı tablosu; masraf paneli. 900 px.',
    'Detaylar butonu gri standart stile alındı.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Seyahat Yönetimi açılınca Talepler sekmesi.',
    'Onay e-postası ?tab=requests → Talepler.',
    'Ekip Seyahatleri eski listeyle aynı.',
    'Giderler’de avanssız+masrafsız seyahat yok.',
    'Masraf bağlayınca Giderler listesinde belirir.',
    'Kart Toplam Masraf = listedeki satır toplamları.',
    'Departman / arama / durum süzgeci kartları da daraltır.',
    'Sayfa 2’ye geçince kartlar değişmez.',
    'Cepten + Şirket = Toplam Masraf.',
    'Ödenen Avans ≤ Toplam Avans.',
    'Detay başlığı “Ad Soyad - Seyahat Detayı”.',
    'Boş alanlar satır açmaz.',
    'Onay Akışı: fotoğraf + ad + unvan + not.',
    'Red adımı kırmızı, gerekçe Not’ta.',
    'Sırası gelmemiş adımlar Beklemede (Bilinmeyen değil).',
    'Talepler satır detayında da masraf paneli.',
    'Üç ekranda Detaylar butonu gri standart.',
    'Altı dilde sekme/kolon/kart çevrili.',
])

# 30
add_section_heading(doc, '30. Avans Yönetimi Ekranı Diğer Modüllerle Aynılaştırıldı', page_break=True)
add_meta(doc, 'Ekran: Bordro Yönetimi → Avans Yönetimi (5 sekme) · Avans Detayı')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Sekmeler: Talepler, Ekip Avansları, Ödeme Bekleyen, Mahsuplaşma, Ek Ödeme Bekleyen. Sayfa Talepler ile açılır.',
    'Durum alt sekmeleri yalnız Talepler ve Ekip Avansları’nda. Kovalar: Bekleyen / Onaylanan / Reddedilen.',
    'Mevcut Onaylayan kolonu; çalışan araması (departman süzgeci yok — bilinçli). ?tab= adresleme.',
    'Detay iç sekmeler kalktı; Alan/Değer + onay + taksit/mahsup alt alta. Sunucu sayfalama 10’ar.',
])
add_subheading(doc, 'Test Maddeleri')
add_check_group(doc, [
    'Her sekmede başlık + açıklama + bilgi butonu.',
    'Sekme sırası doğru; sayfa Talepler ile açılır.',
    'İkon + sayı rozeti; parantezli sayı yok.',
    'Talepler ve Ekip Avansları’nda durum alt sekmeleri + dört kart.',
    'Diğer üç sekmede durum alt sekmesi yok.',
    'Durum alt sekmesi listeyi süzer; kartlar değişmez.',
    'Onaylanan kovada ödenmiş/taksitteki avanslar da var.',
    'İptal → Reddedilen; çalışan onayı bekleyen → Bekleyen.',
    'Satır zeminleri duruma göre renkli.',
    'Mevcut Onaylayan: fotoğraf + ad + unvan; red/iptalde “-”.',
    'İki adımlı akışta kolon ikinci onaylayanı gösterir.',
    'Sayfa açılır açılmaz beş rozet dolu.',
    'Arama Türkçe karakter; Filtreleri Temizle; sekme değişince arama korunur.',
    'Eşleşme yoksa “veri yok”, hata yok.',
    'Yenilemede aynı sekme; ?tab=teamAdvances; geçersiz tab → Talepler.',
    'Gömülü sekmede ikinci sayfa başlığı yok.',
    'Detay iç sekmesiz; boş alan satır açmaz.',
    'Maaş avansında taksitler var, mahsup yok; iş avansında tersi.',
    'Çalışan onayı beklerken sarı uyarı işaretleri durur.',
    'Onay akışı tablosu standart; ödenmemiş maaşta taksit tutarı plandan hesaplanır.',
    'Altı dilde başlık/açıklama/kolon çevrili.',
    'Sayfalama 10 kayıt; 2. sayfa tekrar/atlama yok.',
    '10’dan az kayıtta tek sayfa; arama/durum değişince 1. sayfa.',
    '3. sayfada onay → rozet/kart güncellenir; diğer sekme sayfası etkilenmez.',
    'İlk yükte yalnız Talepler listesi iner.',
    'Mahsuplaşma paketleri sayfa sınırında bölünmez.',
    'Çalışan hücresi ortak kart; fotosuz avatar.',
    'Mevcut Onaylayan artık “-” değil (önceki tur boştu).',
    'Mobil: ad altında durum/tutar; İşlemler sabit.',
])

# 31
add_section_heading(doc, '31. Avans: Toplu İşlem, Hatırlatıcı ve Tıklanabilir Mevcut Onaylayan', page_break=True)
add_meta(doc, 'Ekran: Avans Yönetimi → Talepler (Toplu İşlem) · Ekip Avansları (Hatırlatıcı, Mevcut Onaylayan)')
add_subheading(doc, 'Yenilik Özeti')
add_bullets(doc, [
    'Toplu İşlem: çalışan kırılımı, not, sırayla işleme; son adımda eksik alan onayda atlanır, ret geçer.',
    'Hatırlatıcı: mevcut onaylayan kırılımı; 10 dk sunucu kilidi; e-postasız seçilemez.',
    'Mevcut Onaylayan ortak çalışan kartına geçti; boş pozisyonda “-” (tıklanmaz).',
])
add_h3(doc, 'Toplu İşlem')
add_check_group(doc, [
    'Talepler arama şeridinde Toplu İşlem butonu.',
    'Pencere çalışan kırılımında; rozet talep sayısı.',
    'Satır tıklanınca talepler açılır (tür, tutar, tarih, adım, Detaylar).',
    'Detaylar toplu işlem penceresinin üstünde.',
    'Çalışan tiki tümünü seçer; biri kalkınca yarım tik.',
    'Tümünü seç; süzgeçte görünmeyenlerin seçimi düşer.',
    'Onayla → not penceresi → “N onaylandı”; rozetler güncellenir.',
    'Reddet aynı akış, sarı uyarı.',
    'Son adımda onaylanan tutarı boş maaş avansı onayda atlanır; reddedilebilir.',
    'Taksit sayısı / ilk ay boş son adım da atlanır.',
    'İlk adım toplu onay → ikinci onaylayana düşer.',
    'Sırası kendinde olmayan talep pencerede yok.',
    'Kısmi hata: “N başarısız”; başarılılar geri alınmaz.',
    'Hepsi başarılıysa pencere kapanır; atlanan varsa açık kalır.',
    'Bekleyen yoksa “kayıt yok”.',
])
add_h3(doc, 'Hatırlatıcı / Mevcut Onaylayan / Genel')
add_check_group(doc, [
    'Ekip Avansları’nda Hatırlatıcı butonu.',
    'Onaylayan kırılımı: fotoğraf, ad, unvan, birim, bekleyen sayısı.',
    'Satır tıklanınca talepler + Detaylar.',
    'Hatırlat → gönderildi; hemen tekrar → “zaten gönderildi”.',
    'Sayfa yenile → 10 dk kilidi durur.',
    'Toplu hatırlatma kişi kişi sonuç.',
    'E-postasız onaylayan seçilemez.',
    'Onaylayan süzgeci tek kişiye iner.',
    'Bekleyen yoksa “kayıt yok”.',
    'E-posta butonu Avans Yönetimi’ni açar.',
    'Birim yetkisi dışı talep listede yok / mail gitmez.',
    'Mevcut Onaylayan ada tıklanınca çalışan kartı.',
    'Fotoğraf + ad + unvan; fotosuz avatar; red/iptalde “-”.',
    'Pozisyonda çalışan yoksa “-”; tıklanınca boş pencere yok.',
    'İki adımlı: hücre ikinci onaylayanı gösterir.',
    'Altı dilde pencereler; TR/EN e-posta metni.',
    'Mobil: butonlar etiketli; tablolar kaydırılabilir.',
])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
run = p.add_run(
    'Not: Bölümler 2–31 yeni sayfada başlar. Kutulara tıklayarak işaretleyin. '
    'Masraf (22–26) için AddExpenseManagement; Seyahat (27–29) için AddTravelModule gerekir.'
)
run.font.size = Pt(8)
run.font.italic = True
run.font.color.rgb = RGBColor(110, 110, 110)

output_path = r'C:\Users\ardak\OneDrive\Masaüstü\DHR_Release_Test_Plani_Eylul2026.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
print(f'Checkboxes: {_CB_ID - 1000}')
